#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
EUREKA 2026 — Compilador maestro del libro Excel oficial.

Cruza Firestore (participantes, evaluaciones, paneles de firmas) y genera
RESULTADOS_OFICIALES_EUREKA_UGEL03_2026.xlsx con cinco hojas:

  1. Ganadores por Area      1.er puesto de cada combinacion categoria x area.
                             Para D y E marca CLASIFICA A ETAPA DRE.
  2. Tres Primeros Puestos   Podio completo por categoria y area.
  3. Escrutinio Total        Todas las filas con las columnas estandar.
  4. Medallero por I.E.      Instituciones ordenadas por oro, plata y bronce.
  5. Auditoria de Captura    Por cada evaluacion: participante, casillero de jurado,
                             evaluador operativo, fecha de registro, ediciones y los tres
                             firmantes oficiales que suscribieron el anexo.

La hoja 5 no es opcional: es el registro de trazabilidad del modelo de firma diferida y
lo que permite reconstruir quien ingreso cada nota aunque el anexo lo suscriban otros.

Uso:
    python build_reporte_maestro_eureka.py [--cred RUTA] [--salida RUTA]
                                           [--entrada-json RUTA]

Con --entrada-json se puede alimentar el compilador desde un volcado local en vez de
Firestore, con la forma {"participantes": [...], "evaluaciones": [...], "paneles": {...}}.
"""

import argparse
import json
import os
import sys
from collections import defaultdict

AQUI = os.path.dirname(os.path.abspath(__file__))
CRED_POR_DEFECTO = os.path.normpath(os.path.join(AQUI, "..", "..", "serviceAccountKey.json"))
SALIDA_POR_DEFECTO = "RESULTADOS_OFICIALES_EUREKA_UGEL03_2026.xlsx"

EVENTO = "XXXVI Feria Escolar Nacional de Ciencia y Tecnologia Eureka"
ANIO = 2026
ETAPA = "UGEL"
DRE = "DRE LIMA METROPOLITANA"
UGEL = "UGEL 03"
NUM_JURADOS = 3
CRITERIO_ORDEN = "promedio"   # 'promedio' | 'suma'

CATEGORIAS = {
    "A": {"nombre": "Categoria A", "nivel": "Primaria", "finaliza_en_ugel": True},
    "B": {"nombre": "Categoria B", "nivel": "Primaria", "finaliza_en_ugel": True},
    "C": {"nombre": "Categoria C", "nivel": "Primaria", "finaliza_en_ugel": True},
    "D": {"nombre": "Categoria D", "nivel": "Secundaria", "finaliza_en_ugel": False},
    "E": {"nombre": "Categoria E", "nivel": "Secundaria", "finaliza_en_ugel": False},
}

AREAS = {
    "ind_ciencia_tecnologia": "Indagacion en Ciencia y Tecnologia",
    "indagacion_social": "Indagacion social",
    "indagacion_cientifica": "Indagacion cientifica",
    "soluciones_tecnologicas": "Soluciones tecnologicas",
    "ciencias_sociales": "Ciencias Sociales",
}

ORDEN_AREAS = list(AREAS.keys())


# ─────────────────────────── Lectura de datos ───────────────────────────

def leer_de_firestore(ruta_cred):
    try:
        import firebase_admin
        from firebase_admin import credentials, firestore
    except ImportError:
        print("Falta la dependencia firebase-admin. Instalela con: pip install firebase-admin")
        sys.exit(1)

    if not os.path.exists(ruta_cred):
        print("No se encontro la clave de cuenta de servicio: %s" % ruta_cred)
        sys.exit(1)

    cred = credentials.Certificate(ruta_cred)
    if not firebase_admin._apps:
        firebase_admin.initialize_app(cred)
    db = firestore.client()

    participantes = [dict(d.to_dict(), id=d.id) for d in db.collection("eurekaParticipantes").stream()]
    evaluaciones = [dict(d.to_dict(), id=d.id) for d in db.collection("eurekaEvaluaciones").stream()]
    paneles = {d.id: dict(d.to_dict(), id=d.id) for d in db.collection("eurekaPanelFirmas").stream()}
    return participantes, evaluaciones, paneles


def leer_de_json(ruta):
    with open(ruta, "r", encoding="utf-8") as f:
        data = json.load(f)
    return data.get("participantes", []), data.get("evaluaciones", []), data.get("paneles", {})


# ─────────────────────────── Logica de dominio ───────────────────────────

def resolver_panel(paneles, categoria, area_id):
    """Cascada de resolucion: del mas especifico al mas general. Gana el primero sellado."""
    for scope in ("CAT_%s__AREA_%s" % (categoria, area_id), "CAT_%s" % categoria, "GLOBAL"):
        p = paneles.get(scope)
        if p and p.get("estado") == "sellado":
            return p
    return None


def firmantes_de(panel):
    if not panel:
        return [None] * NUM_JURADOS
    por_slot = {f.get("numeroJurado"): f for f in panel.get("firmantes", [])}
    return [por_slot.get(i + 1) for i in range(NUM_JURADOS)]


def texto(valor, defecto=""):
    return str(valor).strip() if valor not in (None, "") else defecto


def nombres_estudiantes(p):
    partes = []
    for e in (p.get("estudiantes") or []):
        nombre = " ".join(filter(None, [
            texto(e.get("apellidoPaterno")), texto(e.get("apellidoMaterno")), texto(e.get("nombres"))
        ])).strip()
        if nombre:
            dni = texto(e.get("numeroDocumento"))
            partes.append("%s%s" % (nombre, (" (DNI %s)" % dni) if dni else ""))
    return " | ".join(partes)


def construir_filas(participantes, evaluaciones):
    """Filas del consolidado por categoria + area, con puesto asignado."""
    por_clave = defaultdict(dict)
    for ev in evaluaciones:
        slot = (ev.get("jurado") or {}).get("numeroJurado")
        if not slot:
            continue
        por_clave[ev.get("participanteId")][slot] = ev

    filas = []
    for p in participantes:
        evals = por_clave.get(p.get("id"), {})
        notas = []
        registradas = 0
        no_prosigue = False
        for slot in range(1, NUM_JURADOS + 1):
            ev = evals.get(slot)
            if ev and ev.get("estado") == "registrada":
                notas.append(float(ev.get("puntajeTotal") or 0))
                registradas += 1
            else:
                notas.append(None)
            if ev and ev.get("noProsigue"):
                no_prosigue = True

        completo = registradas == NUM_JURADOS
        suma = sum(n for n in notas if n is not None) if completo else None
        promedio = round(suma / NUM_JURADOS, 2) if completo else None
        total = promedio if CRITERIO_ORDEN == "promedio" else suma

        inst = p.get("institucion") or {}
        filas.append({
            "participante_id": p.get("id"),
            "codigo": texto(p.get("codigoParticipante"), p.get("id", "")),
            "categoria": p.get("categoria"),
            "area_id": p.get("areaId"),
            "linea_id": texto(p.get("lineaId")),
            "anexo": texto(p.get("anexoEvaluacion")),
            "institucion": texto(inst.get("nombre"), texto(p.get("institucionNombre"), "I. E. no registrada")),
            "codigo_modular": texto(inst.get("codigoModular")),
            "ugel": texto(inst.get("ugel"), UGEL),
            "dre": texto(inst.get("dre"), DRE),
            "proyecto": texto(p.get("tituloProyecto"), "Sin titulo registrado"),
            "orden": p.get("ordenPresentacion") or 0,
            "notas": notas,
            "suma": suma,
            "promedio": promedio,
            "total": total,
            "completo": completo,
            "no_prosigue": no_prosigue,
            "no_se_presento": bool(p.get("noSePresento")),
            "orden_manual": p.get("ordenManual"),
            "estudiantes": nombres_estudiantes(p),
            "docente": texto((p.get("docenteAsesor") or {}).get("nombreCompleto")),
            "docente_dni": texto((p.get("docenteAsesor") or {}).get("numeroDocumento")),
            "docente_especialidad": texto((p.get("docenteAsesor") or {}).get("especialidad")),
            "docente_telefono": texto((p.get("docenteAsesor") or {}).get("telefono")),
            "docente_correo": texto((p.get("docenteAsesor") or {}).get("correo")),
            "enlace": texto(p.get("urlTrabajo")),
            "puesto": None,
        })

    # Puestos por combinacion categoria + area, solo sobre filas elegibles.
    grupos = defaultdict(list)
    for f in filas:
        grupos[(f["categoria"], f["area_id"])].append(f)

    for elementos in grupos.values():
        elegibles = [
            f for f in elementos
            if f["completo"] and not f["no_se_presento"] and not f["no_prosigue"]
            and f["total"] is not None and f["total"] > 0
        ]
        elegibles.sort(key=lambda f: (
            -f["total"],
            f["orden_manual"] if f["orden_manual"] is not None else 9999,
            f["orden"]
        ))
        puesto = 1
        for i, f in enumerate(elegibles):
            if i > 0:
                prev = elegibles[i - 1]
                empatan = (f["total"] == prev["total"]
                           and f["orden_manual"] is None and prev["orden_manual"] is None)
                if not empatan:
                    puesto = i + 1
            f["puesto"] = puesto

    return filas


# ─────────────────────────── Construccion del libro ───────────────────────────

def construir_libro(filas, evaluaciones, participantes, paneles, salida):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    VERDE = "6E9E23"
    NAVY = "122240"

    cabecera_font = Font(bold=True, color="FFFFFF", size=10)
    cabecera_fill = PatternFill("solid", fgColor=VERDE)
    titulo_font = Font(bold=True, size=13, color=NAVY)
    borde = Border(*(Side(style="thin", color="D1D5DB"),) * 4)

    wb = Workbook()

    def nueva_hoja(titulo, encabezados, primera=False):
        ws = wb.active if primera else wb.create_sheet()
        ws.title = titulo[:31]
        ws["A1"] = "%s %d" % (EVENTO, ANIO)
        ws["A1"].font = titulo_font
        ws["A2"] = "Etapa %s  -  %s  -  %s" % (ETAPA, DRE, UGEL)
        ws["A2"].font = Font(size=9, color="6B7280")
        ws["A3"] = titulo
        ws["A3"].font = Font(bold=True, size=11, color=NAVY)
        for col, texto_col in enumerate(encabezados, start=1):
            c = ws.cell(row=5, column=col, value=texto_col)
            c.font = cabecera_font
            c.fill = cabecera_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = borde
        ws.freeze_panes = "A6"
        return ws

    def autoajustar(ws, ancho_max=52):
        for col in ws.columns:
            largo = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, largo + 2), ancho_max)

    def nombre_area(area_id):
        return AREAS.get(area_id, area_id)

    # ── Hoja 1: Ganadores por Area ──
    ws1 = nueva_hoja("1. Ganadores por Area", [
        "Categoria", "Nivel", "Area de participacion", "I. E.", "Codigo Modular",
        "Proyecto", "Puntaje total", "Promedio", "Estudiantes", "Docente asesor", "Observacion"
    ], primera=True)

    fila = 6
    for cat in sorted(CATEGORIAS):
        for area_id in ORDEN_AREAS:
            ganador = next((f for f in filas
                            if f["categoria"] == cat and f["area_id"] == area_id and f["puesto"] == 1), None)
            if not ganador:
                continue
            observacion = ("CLASIFICA A ETAPA DRE" if not CATEGORIAS[cat]["finaliza_en_ugel"]
                           else "La participacion finaliza en la etapa UGEL")
            valores = [
                CATEGORIAS[cat]["nombre"], CATEGORIAS[cat]["nivel"], nombre_area(area_id),
                ganador["institucion"], ganador["codigo_modular"], ganador["proyecto"],
                ganador["suma"], ganador["promedio"], ganador["estudiantes"], ganador["docente"], observacion
            ]
            for col, v in enumerate(valores, start=1):
                c = ws1.cell(row=fila, column=col, value=v)
                c.border = borde
                if col == 11 and observacion.startswith("CLASIFICA"):
                    c.font = Font(bold=True, color="1E4D7B")
            fila += 1
    autoajustar(ws1)

    # ── Hoja 2: Tres Primeros Puestos ──
    ws2 = nueva_hoja("2. Tres Primeros Puestos", [
        "Categoria", "Area de participacion", "Puesto", "I. E.", "Proyecto",
        "Puntaje total", "Promedio", "Estudiantes", "Docente asesor"
    ])
    fila = 6
    for cat in sorted(CATEGORIAS):
        for area_id in ORDEN_AREAS:
            podio = sorted(
                [f for f in filas if f["categoria"] == cat and f["area_id"] == area_id
                 and f["puesto"] and f["puesto"] <= 3],
                key=lambda f: f["puesto"]
            )
            for f in podio:
                valores = [
                    CATEGORIAS[cat]["nombre"], nombre_area(area_id), "%d." % f["puesto"],
                    f["institucion"], f["proyecto"], f["suma"], f["promedio"],
                    f["estudiantes"], f["docente"]
                ]
                for col, v in enumerate(valores, start=1):
                    c = ws2.cell(row=fila, column=col, value=v)
                    c.border = borde
                fila += 1
    autoajustar(ws2)

    # ── Hoja 3: Escrutinio Total ──
    encabezados3 = [
        "Puesto", "Promedio", "Suma", "Jurado 1", "Jurado 2", "Jurado 3",
        "I. E.", "Codigo Modular", "Codigo SICE", "Proyecto", "Categoria",
        "Area de participacion", "Linea", "Anexo", "Estudiantes",
        "Docente asesor", "DNI docente", "Especialidad", "Telefono", "Correo",
        "Enlace a la evidencia", "Completo", "No prosigue", "No se presento", "Orden de presentacion"
    ]
    ws3 = nueva_hoja("3. Escrutinio Total", encabezados3)
    fila = 6
    ordenadas = sorted(filas, key=lambda f: (
        f["categoria"] or "",
        ORDEN_AREAS.index(f["area_id"]) if f["area_id"] in ORDEN_AREAS else 99,
        f["puesto"] if f["puesto"] else 999,
        f["orden"]
    ))
    for f in ordenadas:
        valores = [
            ("%d." % f["puesto"]) if f["puesto"] else "-",
            f["promedio"], f["suma"],
            f["notas"][0], f["notas"][1], f["notas"][2],
            f["institucion"], f["codigo_modular"], f["codigo"], f["proyecto"],
            CATEGORIAS.get(f["categoria"], {}).get("nombre", f["categoria"]),
            nombre_area(f["area_id"]), f["linea_id"], f["anexo"], f["estudiantes"],
            f["docente"], f["docente_dni"], f["docente_especialidad"],
            f["docente_telefono"], f["docente_correo"], f["enlace"],
            "Si" if f["completo"] else "No",
            "Si" if f["no_prosigue"] else "No",
            "Si" if f["no_se_presento"] else "No",
            f["orden"]
        ]
        for col, v in enumerate(valores, start=1):
            c = ws3.cell(row=fila, column=col, value=v)
            c.border = borde
        fila += 1
    autoajustar(ws3, ancho_max=44)

    # ── Hoja 4: Medallero por I.E. ──
    medallero = defaultdict(lambda: {"oro": 0, "plata": 0, "bronce": 0, "proyectos": 0})
    for f in filas:
        m = medallero[f["institucion"]]
        m["proyectos"] += 1
        if f["puesto"] == 1:
            m["oro"] += 1
        elif f["puesto"] == 2:
            m["plata"] += 1
        elif f["puesto"] == 3:
            m["bronce"] += 1

    ws4 = nueva_hoja("4. Medallero por I.E.", [
        "I. E.", "Primeros puestos", "Segundos puestos", "Terceros puestos",
        "Total en el podio", "Proyectos presentados"
    ])
    fila = 6
    for inst, m in sorted(medallero.items(), key=lambda kv: (-kv[1]["oro"], -kv[1]["plata"], -kv[1]["bronce"], kv[0])):
        valores = [inst, m["oro"], m["plata"], m["bronce"], m["oro"] + m["plata"] + m["bronce"], m["proyectos"]]
        for col, v in enumerate(valores, start=1):
            c = ws4.cell(row=fila, column=col, value=v)
            c.border = borde
        fila += 1
    autoajustar(ws4)

    # ── Hoja 5: Auditoria de Captura ──
    ws5 = nueva_hoja("5. Auditoria de Captura", [
        "Codigo SICE", "I. E.", "Proyecto", "Categoria", "Area de participacion",
        "Casillero de jurado", "Estado", "Puntaje total",
        "Evaluador operativo", "Correo del evaluador", "Registrada el",
        "Ediciones posteriores", "Ultima edicion",
        "Panel que gobierna", "Estado del panel", "Firmante 1", "Firmante 2", "Firmante 3"
    ])

    por_id = {p.get("id"): p for p in participantes}
    fila = 6
    for ev in sorted(evaluaciones, key=lambda e: (
        e.get("categoria") or "",
        ORDEN_AREAS.index(e.get("areaId")) if e.get("areaId") in ORDEN_AREAS else 99,
        e.get("participanteId") or "",
        (e.get("jurado") or {}).get("numeroJurado") or 0
    )):
        p = por_id.get(ev.get("participanteId"), {})
        inst = (p.get("institucion") or {})
        panel = resolver_panel(paneles, ev.get("categoria"), ev.get("areaId"))
        firmantes = firmantes_de(panel)
        operativo = ev.get("evaluadorOperativo") or {}
        historial = ev.get("historialEdiciones") or []

        valores = [
            texto(p.get("codigoParticipante"), texto(ev.get("participanteId"))),
            texto(inst.get("nombre"), texto(p.get("institucionNombre"), "I. E. no registrada")),
            texto(p.get("tituloProyecto"), "Sin titulo registrado"),
            CATEGORIAS.get(ev.get("categoria"), {}).get("nombre", ev.get("categoria")),
            nombre_area(ev.get("areaId")),
            "Jurado N. %s" % ((ev.get("jurado") or {}).get("numeroJurado")),
            texto(ev.get("estado")),
            ev.get("puntajeTotal"),
            texto(operativo.get("nombreCompleto")),
            texto(operativo.get("correo")),
            texto(ev.get("registradaEn")),
            len(historial),
            texto(historial[-1].get("en")) if historial else "",
            panel.get("id") if panel else "Sin panel sellado",
            panel.get("estado") if panel else "preliminar",
            texto((firmantes[0] or {}).get("nombreCompleto")),
            texto((firmantes[1] or {}).get("nombreCompleto")),
            texto((firmantes[2] or {}).get("nombreCompleto")),
        ]
        for col, v in enumerate(valores, start=1):
            c = ws5.cell(row=fila, column=col, value=v)
            c.border = borde
        fila += 1
    autoajustar(ws5, ancho_max=40)

    wb.save(salida)
    return {
        "hojas": len(wb.sheetnames),
        "filas": len(filas),
        "evaluaciones": len(evaluaciones),
    }


def main():
    parser = argparse.ArgumentParser(description="Compila el libro Excel oficial de Eureka 2026.")
    parser.add_argument("--cred", default=CRED_POR_DEFECTO, help="Clave de cuenta de servicio de Firebase")
    parser.add_argument("--salida", default=SALIDA_POR_DEFECTO, help="Ruta del archivo .xlsx de salida")
    parser.add_argument("--entrada-json", default=None,
                        help="Volcado local con {participantes, evaluaciones, paneles} en vez de Firestore")
    args = parser.parse_args()

    if args.entrada_json:
        participantes, evaluaciones, paneles = leer_de_json(args.entrada_json)
        print("Origen: volcado local %s" % args.entrada_json)
    else:
        participantes, evaluaciones, paneles = leer_de_firestore(args.cred)
        print("Origen: Firestore")

    print("Participantes: %d  |  Evaluaciones: %d  |  Paneles: %d"
          % (len(participantes), len(evaluaciones), len(paneles)))

    if not participantes:
        print("No hay participantes registrados. No se genera el libro.")
        return 1

    filas = construir_filas(participantes, evaluaciones)
    resumen = construir_libro(filas, evaluaciones, participantes, paneles, args.salida)

    print("")
    print("Libro generado: %s" % os.path.abspath(args.salida))
    print("Hojas: %d  |  Filas de escrutinio: %d  |  Evaluaciones auditadas: %d"
          % (resumen["hojas"], resumen["filas"], resumen["evaluaciones"]))

    sin_panel = [e for e in evaluaciones
                 if not resolver_panel(paneles, e.get("categoria"), e.get("areaId"))]
    if sin_panel:
        print("")
        print("ATENCION: %d evaluacion(es) no tienen un panel de firmas sellado aplicable."
              % len(sin_panel))
        print("Sus anexos se emiten como documentos preliminares sin valor legal.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
